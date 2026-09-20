"""Text extraction from uploaded documents, preserving page numbers.

Page numbers are the citation unit people actually expect from a document
platform ("page 5", not "chunk 14"), so extraction returns a list of
``(page_number, text)`` pairs rather than one flat string - chunking then
carries the page number through to every chunk it produces.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt", ".md"}


@dataclass
class ExtractedPage:
    page_number: int  # 1-indexed; always 1 for page-less formats
    text: str


@dataclass
class ExtractedDocument:
    pages: list[ExtractedPage]
    page_count: int

    @property
    def full_text(self) -> str:
        return "\n\n".join(p.text for p in self.pages if p.text.strip())


def extract_pdf(data: bytes) -> ExtractedDocument:
    import io

    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = [
        ExtractedPage(page_number=i + 1, text=(page.extract_text() or ""))
        for i, page in enumerate(reader.pages)
    ]
    return ExtractedDocument(pages=pages, page_count=len(pages))


def extract_docx(data: bytes) -> ExtractedDocument:
    import io

    from docx import Document as DocxDocument

    doc = DocxDocument(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    return ExtractedDocument(pages=[ExtractedPage(page_number=1, text=text)], page_count=1)


def extract_plain_text(data: bytes) -> ExtractedDocument:
    text = data.decode("utf-8", errors="replace")
    return ExtractedDocument(pages=[ExtractedPage(page_number=1, text=text)], page_count=1)


def extract(data: bytes, suffix: str) -> ExtractedDocument:
    suffix = suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(data)
    if suffix == ".docx":
        return extract_docx(data)
    if suffix in {".txt", ".md"}:
        return extract_plain_text(data)
    raise ValueError(f"Unsupported file type: {suffix}")


def extract_file(path: Path) -> ExtractedDocument:
    return extract(path.read_bytes(), path.suffix)
